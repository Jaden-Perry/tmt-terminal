"""Status, universe, search, watchlists, dashboard, export."""
from __future__ import annotations

import io
import time

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from .. import config, db, ratelimit
from ..services import dashboard, refresh, universe, watchlists

router = APIRouter(prefix="/api")


@router.get("/status")
def status():
    c = db.conn()
    n_prices = c.execute("SELECT COUNT(DISTINCT ticker) FROM prices").fetchone()[0]
    n_metrics = c.execute("SELECT COUNT(*) FROM metrics").fetchone()[0]
    jobs = [dict(r) for r in c.execute("SELECT job, started, finished, status, substr(detail,1,300) AS detail FROM job_log ORDER BY id DESC LIMIT 8")]
    last_price = db.cache_get("prices:last_refresh")
    return {
        "keys": config.key_status(),
        "universe_count": len(universe.tickers()),
        "tickers_with_prices": n_prices,
        "tickers_with_snapshots": n_metrics,
        "last_price_refresh": last_price[0]["at"] if last_price else None,
        "refresh": refresh.STATUS,
        "jobs": jobs,
        "api_calls_today": {s: ratelimit.calls_today(s) for s in ("fmp", "finnhub", "fred", "anthropic")},
        "fmp_daily_budget": config.FMP_DAILY_BUDGET,
    }


@router.post("/refresh/{job}")
def trigger(job: str):
    fns = {"daily": refresh.daily_quick, "nightly": refresh.nightly, "recompute": refresh.recompute_local,
           "universe": lambda: universe.build_universe()}
    if job not in fns:
        raise HTTPException(404, "unknown job")
    started = refresh.run_in_background(job, fns[job])
    return {"started": started, "message": None if started else "Another refresh is already running."}


@router.get("/universe")
def get_universe():
    return {"rows": universe.all_rows(), "sub_industries": universe.SUB_INDUSTRIES}


@router.get("/search")
def search(q: str):
    q = q.strip().upper()
    if not q:
        return []
    rows = universe.all_rows()
    hits = [r for r in rows if r["ticker"].startswith(q)]
    hits += [r for r in rows if r not in hits and q in (r["name"] or "").upper()]
    out = [{"ticker": r["ticker"], "name": r["name"], "sub_industry": r["sub_industry"]} for r in hits[:12]]
    if not any(o["ticker"] == q for o in out):
        cmap = db.cache_get("sec:cik_map")
        if cmap and q in cmap[0]:
            out.insert(0, {"ticker": q, "name": "(outside TMT universe)", "sub_industry": None})
    return out


@router.get("/watchlists")
def get_watchlists():
    return watchlists.all_lists()


class WLItem(BaseModel):
    name: str = watchlists.DEFAULT
    ticker: str


@router.post("/watchlists")
def add_watch(item: WLItem):
    watchlists.add(item.name, item.ticker)
    return watchlists.all_lists()


@router.delete("/watchlists/{name}/{ticker}")
def remove_watch(name: str, ticker: str):
    watchlists.remove(name, ticker)
    return watchlists.all_lists()


@router.get("/dashboard/watchlist")
def dash_watchlist(name: str = watchlists.DEFAULT):
    return {"name": name, "rows": dashboard.watchlist_table(name), "generated_at": time.time()}


@router.get("/dashboard/heatmap")
def dash_heatmap(period: str = "1d"):
    return dashboard.heatmap(period)


@router.get("/dashboard/movers")
def dash_movers(period: str = "1d", n: int = 10):
    return dashboard.movers(period, n)


@router.get("/dashboard/benchmarks")
def dash_benchmarks():
    return dashboard.benchmarks()


class Table(BaseModel):
    title: str = "export"
    columns: list[str]
    rows: list[list]


@router.post("/export/xlsx")
def export_xlsx(t: Table):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (t.title or "export")[:31]
    ws.append(t.columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
    for r in t.rows:
        ws.append([None if v in ("N/A", "") else v for v in r])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in t.title) + ".xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/screen")
def screen():
    from ..services import screener

    return screener.run()


@router.get("/news")
def news_feed(scope: str = "sector", watchlist: str = watchlists.DEFAULT, ticker: str | None = None,
              sub_industry: str | None = None, since: str | None = None, until: str | None = None,
              deals: bool = False, q: str | None = None, refresh: bool = False):
    from ..services import news

    notes: list = []
    tickers = None
    if ticker:
        tickers = [ticker.upper()]
        news.refresh_company(ticker.upper())
    elif scope == "watchlist":
        tickers = watchlists.all_lists().get(watchlist, [])
        for t in tickers:
            news.refresh_company(t)
    else:
        notes = news.refresh_sector(ttl=0 if refresh else 3 * db.HOUR)
    items = news.query(tickers=tickers, sub_industry=sub_industry, since=since, until=until, deals_only=deals, q=q, limit=400)
    if scope == "sector" and not ticker and not sub_industry:
        items = [i for i in items if i["ticker"] is None or universe.get(i["ticker"])]
    return {"items": items, "notes": notes, "count": len(items)}


@router.get("/ai/usage")
def ai_usage():
    from ..services import ai

    return ai.month_spend()


@router.get("/macro")
def macro_dashboard(years: int = 10):
    from ..services import macro

    return macro.dashboard(years)


@router.get("/macro/sensitivity")
def macro_sensitivity(watchlist: str = watchlists.DEFAULT, years: float = 3):
    from ..services import macro

    return macro.watchlist_sensitivities(watchlist, years)
